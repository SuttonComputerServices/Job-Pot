import io

import pytest
from docx import Document
from pypdf import PdfWriter

from jobpot.db import SessionLocal
from jobpot.models import Criteria, Resume
from jobpot.resume import extract_text, review

from .helpers import client, csrf, login, make_user

WEAK = """Jane Doe
Summary
I was responsible for many things and I helped with projects. My duties included reports and I worked on a lot of different tasks across the whole company, which meant that I was regularly involved in meetings, planning sessions, budgets, hiring and many other kinds of work for my team.
"""

STRONG_BULLETS = "\n".join(
    f"- Led migration {i} of billing services to Python, cutting monthly costs by {10 + i}% and saving ${i}0,000 a year"
    for i in range(1, 25))
STRONG = f"""Jane Doe
jane@example.com | +1 555 123 4567 | linkedin.com/in/janedoe
Experience
Senior Python developer, Acme, 2019 to 2024
{STRONG_BULLETS}
Education
BSc Computer Science, State University
Skills
Python, SQL, FastAPI, Docker
"""


def _docx(paragraphs, table=None) -> bytes:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    doc.add_paragraph("Shipped the API", style="List Bullet")
    if table:
        t = doc.add_table(rows=1, cols=len(table))
        for cell, value in zip(t.rows[0].cells, table):
            cell.text = value
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _blank_pdf() -> bytes:
    w = PdfWriter()
    w.add_blank_page(612, 792)
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


def _titles(found, severity):
    return {s.title for s in found if s.severity == severity}


def test_extract_docx_includes_tables_and_bullets():
    text = extract_text("cv.DOCX", _docx(["Jane Doe", "Experience"], table=["Python", "SQL"]))
    assert "Jane Doe" in text and "Python | SQL" in text and "- Shipped the API" in text


def test_extract_txt_with_fallback_encoding():
    assert extract_text("cv.txt", "Résumé".encode()) == "Résumé"
    assert extract_text("cv.md", "Café".encode("cp1252")) == "Café"


def test_unsupported_and_broken_files_raise_value_error():
    with pytest.raises(ValueError, match="PDF, Word"):
        extract_text("cv.doc", b"anything")
    with pytest.raises(ValueError, match="couldn't be read"):
        extract_text("cv.docx", b"not a zip")


def test_blank_pdf_gives_empty_text():
    assert extract_text("scan.pdf", _blank_pdf()) == ""


def test_each_rule_fires_on_weak_sample():
    crit = Criteria(keywords="Kubernetes, Python", titles="Platform engineer")
    found = review(WEAK, crit)
    fix, improve = _titles(found, "fix"), _titles(found, "improve")
    assert {"Add your contact details", "Add more detail", "Use bullet points", "Add your target keywords", "Add section headings"} <= fix
    assert {"Replace weak phrases", "Drop first-person words", "Break up long paragraphs"} <= improve
    keywords = next(s for s in found if s.title == "Add your target keywords")
    assert "Kubernetes" in keywords.detail and "Platform engineer" in keywords.detail


def test_numbers_and_length_rules():
    few_numbers = "Experience\n" + "\n".join("- Built internal tools for the support team" for _ in range(10))
    assert "Show results with numbers" in _titles(review(few_numbers, None), "improve")
    assert "Cut it down" in _titles(review("word " * 1200, None), "improve")


def test_strong_sample_has_fewer_fixes():
    crit = Criteria(keywords="Python, SQL", titles="Python developer")
    strong = review(STRONG, crit)
    assert len(_titles(strong, "fix")) < len(_titles(review(WEAK, crit), "fix"))
    assert not _titles(strong, "fix"), _titles(strong, "fix")
    assert "Set your search criteria" in _titles(review(STRONG, None), "improve")


def test_date_range_is_not_a_phone_number():
    found = review("jane@example.com linkedin.com/in/jane\nAcme 2018 - 2022", None)
    contact = next(s for s in found if s.title == "Add your contact details")
    assert "phone" in contact.detail


def _upload(c, name, data):
    return c.post("/resume/upload", data={"csrf": csrf(c, "/resume")}, files={"file": (name, data)}, follow_redirects=True)


def test_upload_too_large_rejected():
    make_user()
    with client() as c:
        login(c)
        r = _upload(c, "cv.txt", b"a" * (5 * 1024 * 1024 + 1))
        assert r.status_code == 400 and "over 5 MB" in r.text
    with SessionLocal() as db:
        assert db.query(Resume).count() == 0


def test_upload_shows_suggestions_edit_and_download():
    uid = make_user()
    with SessionLocal() as db:
        db.add(Criteria(user_id=uid, keywords="Kubernetes"))
        db.commit()
    with client() as c:
        login(c)
        r = _upload(c, "../My CV.txt", WEAK.encode())
        assert r.status_code == 200
        assert "Fix these first" in r.text and "Kubernetes" in r.text and "My CV.txt" in r.text
        assert "Get an AI review" not in r.text  # AI is off in tests
        d = c.get("/resume/download")
        assert d.content == WEAK.encode() and d.headers["content-disposition"] == 'attachment; filename="My_CV.txt"'
        c.post("/resume/text", data={"text": "Edited text", "csrf": csrf(c, "/resume")})
        t = c.get("/resume/download.txt")
        assert t.text == "Edited text" and "attachment" in t.headers["content-disposition"]
        r = c.post("/resume/ai", data={"csrf": csrf(c, "/resume")})
        assert "off or unreachable" in r.text
    with SessionLocal() as db:
        assert db.query(Resume).count() == 1


def test_other_user_cannot_see_resume():
    make_user()
    make_user("bob")
    with client() as c:
        login(c)
        _upload(c, "cv.txt", b"Secret Alice resume text")
    with client() as c:
        login(c, "bob")
        assert "Secret Alice" not in c.get("/resume").text
        assert c.get("/resume/download").status_code == 404
        assert c.get("/resume/download.txt").status_code == 404


def test_scanned_pdf_shows_scan_message():
    make_user()
    with client() as c:
        login(c)
        r = _upload(c, "scan.pdf", _blank_pdf())
        assert "looks like a scan" in r.text
        assert c.get("/resume/download").headers["content-type"] == "application/pdf"


def test_ai_review_shown_when_ai_on(monkeypatch):
    from jobpot import ai
    prompts = []
    monkeypatch.setattr(ai, "ai_enabled", lambda db: True)
    monkeypatch.setattr(ai, "generate", lambda db, prompt, timeout=60: prompts.append(prompt) or "1. Add numbers <b>")
    make_user()
    with client() as c:
        login(c)
        _upload(c, "cv.txt", STRONG.encode())
        assert "Get an AI review" in c.get("/resume").text
        r = c.post("/resume/ai", data={"csrf": csrf(c, "/resume")})
        assert "1. Add numbers &lt;b&gt;" in r.text and "jane@example.com" in prompts[0]
